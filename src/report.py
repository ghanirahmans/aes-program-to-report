import re
import shutil
import subprocess
import zipfile


try:
    import docx
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    docx = Inches = Pt = WD_ALIGN_PARAGRAPH = WD_TABLE_ALIGNMENT = RGBColor = WD_SECTION = OxmlElement = qn = None


def add_bold_paragraph(doc, text):
    if doc:
        p = doc.add_paragraph()
        p.add_run(text).bold = True


def set_section_narrow_margins(section):
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)


def set_narrow_margins(doc):
    for section in doc.sections:
        set_section_narrow_margins(section)


def set_section_columns(section, column_count):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    cols = cols[0] if cols else OxmlElement("w:cols")
    if cols.getparent() is None:
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(column_count))
    cols.set(qn("w:equalWidth"), "1")


def start_two_column_section(doc):
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_section_narrow_margins(section)
    set_section_columns(section, 2)


def format_compact_matrix_table(table):
    cell_width = Inches(0.5)
    width_twips = str(cell_width.twips)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(cell_width.twips * 4))

    for grid_col in table._tbl.tblGrid.iterchildren():
        grid_col.set(qn("w:w"), width_twips)

    for row in table.rows:
        for cell in row.cells:
            cell.width = cell_width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), width_twips)


def clear_document_metadata(doc):
    core_properties = doc.core_properties
    for attr in (
        "author",
        "category",
        "comments",
        "content_status",
        "identifier",
        "keywords",
        "language",
        "last_modified_by",
        "subject",
        "title",
        "version",
    ):
        try:
            setattr(core_properties, attr, "")
        except AttributeError:
            pass
    core_properties.revision = 1


def strip_docx_metadata(docx_path):
    empty_core_properties = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<cp:coreProperties '
        'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" '
        'xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"/>'
    ).encode("utf-8")
    empty_app_properties = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties '
        'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
    ).encode("utf-8")

    temp_path = docx_path.with_suffix(".tmp.docx")
    fixed_timestamp = (1980, 1, 1, 0, 0, 0)

    with zipfile.ZipFile(docx_path, "r") as source:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "docProps/core.xml":
                    data = empty_core_properties
                elif item.filename == "docProps/app.xml":
                    data = empty_app_properties

                clean_item = zipfile.ZipInfo(item.filename, fixed_timestamp)
                clean_item.compress_type = zipfile.ZIP_DEFLATED
                clean_item.external_attr = item.external_attr
                target.writestr(clean_item, data)

    temp_path.replace(docx_path)


def convert_docx_to_pdf(docx_path, output_dir):
    converter = shutil.which("libreoffice") or shutil.which("soffice")
    if not converter:
        raise RuntimeError("LibreOffice tidak ditemukan, PDF tidak bisa dibuat otomatis.")

    subprocess.run(
        [
            converter,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    return docx_path.with_suffix(".pdf")


def add_calculation_paragraph(doc, text, indent=False):
    if doc:
        p = doc.add_paragraph()
        if indent:
            p.paragraph_format.left_indent = Pt(18)
        run = p.add_run(text.strip())
        font = run.font
        font.name = "Times New Roman"
        font.size = Pt(12)
        font.color.rgb = RGBColor(0, 0, 0)


def add_poly_text_to_paragraph(p, text):
    parts = re.split(r"(x\^\d+)", text)
    for part in parts:
        if part.startswith("x^"):
            base, exponent = part.split("^")
            run = p.add_run(base)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run_sup = p.add_run(exponent)
            run_sup.font.superscript = True
            run_sup.font.name = "Times New Roman"
            run_sup.font.size = Pt(12)
            run_sup.font.color.rgb = RGBColor(0, 0, 0)
        else:
            run = p.add_run(part)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)


def add_single_term_to_para(p, term, is_duplicate):
    font_attrs = {"name": "Times New Roman", "size": Pt(12), "color": {"rgb": RGBColor(0, 0, 0)}}

    if "^" in term:
        base, exponent = term.split("^", 1)
        run = p.add_run(base)
        run.font.name = font_attrs["name"]
        run.font.size = font_attrs["size"]
        run.font.color.rgb = font_attrs["color"]["rgb"]
        run.font.strike = is_duplicate
        run.font.bold = is_duplicate

        run_sup = p.add_run(exponent)
        run_sup.font.superscript = True
        run_sup.font.name = font_attrs["name"]
        run_sup.font.size = font_attrs["size"]
        run_sup.font.color.rgb = font_attrs["color"]["rgb"]
        run_sup.font.strike = is_duplicate
        run_sup.font.bold = is_duplicate
    else:
        run = p.add_run(term)
        run.font.name = font_attrs["name"]
        run.font.size = font_attrs["size"]
        run.font.color.rgb = font_attrs["color"]["rgb"]
        run.font.strike = is_duplicate
        run.font.bold = is_duplicate


def add_matrix_to_doc(doc, label, hex_string):
    if doc:
        p = doc.add_paragraph()
        p.add_run(label).italic = True
        font = p.runs[0].font
        font.name = "Times New Roman"
        font.size = Pt(12)
        table = doc.add_table(rows=4, cols=4)
        table.style = "Table Grid"
        format_compact_matrix_table(table)
        bytes_array = [hex_string[i : i + 2] for i in range(0, len(hex_string), 2)]
        for r in range(4):
            for c in range(4):
                cell = table.cell(r, c)
                run = cell.paragraphs[0].add_run(bytes_array[c * 4 + r].upper())
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0, 0, 0)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()


def add_constant_matrix_to_doc(doc, label, matrix):
    if doc:
        p = doc.add_paragraph()
        p.add_run(label).italic = True
        font = p.runs[0].font
        font.name = "Times New Roman"
        font.size = Pt(12)
        table = doc.add_table(rows=4, cols=4)
        table.style = "Table Grid"
        format_compact_matrix_table(table)
        for r in range(4):
            for c in range(4):
                cell = table.cell(r, c)
                run = cell.paragraphs[0].add_run(f"{matrix[r][c]:02X}")
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0, 0, 0)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()
